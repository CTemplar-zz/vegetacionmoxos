"""Convierte IDZonas.xlsx en un índice compacto de inundación por fecha."""

from __future__ import annotations

import argparse
import base64
import json
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


def iso_date(value) -> str:
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    return str(value)[:10]


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    workbook = load_workbook(args.source, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    headers = list(next(rows))
    date_headers = headers[4:-2]
    dates = [iso_date(value) for value in date_headers]
    row_count = sheet.max_row - 1
    bitsets = [bytearray((row_count + 7) // 8) for _ in dates]
    ids: list[int] = []
    classes: list[str] = []
    fractional_values = 0

    for row_index, row in enumerate(rows):
        segment_id = int(row[0])
        ids.append(segment_id)
        classes.append(str(row[1]))
        for date_index, value in enumerate(row[4:-2]):
            if isinstance(value, (int, float)) and value > 0:
                bitsets[date_index][row_index >> 3] |= 1 << (row_index & 7)
                if value != 1:
                    fractional_values += 1
            elif value not in (None, 0, False):
                raise ValueError(f"Valor no numérico en fila {row_index + 2}, fecha {dates[date_index]}: {value!r}")

    if len(ids) != row_count or len(set(ids)) != len(ids):
        raise ValueError("OBJECTID no es único o el número de filas no coincide.")
    payload = {
        "version": 1,
        "idField": "OBJECTID_1",
        "sourceIdField": "OBJECTID",
        "dates": dates,
        "ids": ids,
        "classes": classes,
        "bitsets": [base64.b64encode(bits).decode("ascii") for bits in bitsets],
        "floodRule": "value > 0",
        "fractionalPositiveValues": fractional_values,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(json.dumps({"segments": len(ids), "dates": len(dates), "bytes": args.output.stat().st_size}))


if __name__ == "__main__":
    main()
